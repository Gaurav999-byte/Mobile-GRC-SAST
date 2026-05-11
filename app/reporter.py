import os
import base64
from jinja2 import Environment, FileSystemLoader
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from xhtml2pdf import pisa


BASE_DIR = os.path.dirname(__file__)
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")

OUT_DIR = os.path.join(os.path.dirname(BASE_DIR), "outputs", "reports")
os.makedirs(OUT_DIR, exist_ok=True)


# =========================
# IMAGE TO BASE64
# =========================
def image_to_base64(image_path):
    with open(image_path, "rb") as img:
        return base64.b64encode(img.read()).decode("utf-8")


# =========================
# CHART GENERATION
# =========================
def _generate_chart(findings, out_dir, base_name):

    counts = {"Critical": 0, "High": 0, "Medium": 0, "Low": 0}

    for f in findings:
        sev = f.get("severity", "Medium")
        if sev not in counts:
            sev = "Medium"
        counts[sev] += 1

    labels = ["Critical", "High", "Medium", "Low"]
    values = [counts[l] for l in labels]

    colors = ["#8B0000", "#FF4500", "#FAD60C", "#32CD32"]

    chart_path = os.path.join(out_dir, f"{base_name}_chart.png")

    plt.figure(figsize=(7, 7), dpi=150)

    plt.pie(
        values,
        labels=labels,
        autopct="%1.0f%%",
        colors=colors,
        startangle=140
    )

    plt.title("Severity Distribution", fontsize=16, fontweight="bold")

    plt.tight_layout()
    plt.savefig(chart_path)
    plt.close()

    return chart_path


# =========================
# MAIN REPORT FUNCTION
# =========================
def render_and_save_report(mapped_findings, filename, out_dir=None):

    out_dir = out_dir or OUT_DIR
    os.makedirs(out_dir, exist_ok=True)

    env = Environment(loader=FileSystemLoader(TEMPLATES_DIR))

    now = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = filename.replace(" ", "_")
    base = f"report_{safe_name}_{now}"

    # =========================
    # CHART
    # =========================
    chart_path = _generate_chart(mapped_findings, out_dir, base)
    chart_base64 = image_to_base64(chart_path)

    # =========================
    # SUMMARY
    # =========================
    summary = {"Critical": 0, "High": 0, "Medium": 0, "Low": 0}

    for f in mapped_findings:
        sev = f.get("severity", "Medium")
        if sev in summary:
            summary[sev] += 1

    # =========================
    # HTML REPORT
    # =========================
    html_template = env.get_template("report_preview.html")

    html_output = html_template.render(
        app_name=safe_name,
        findings=mapped_findings,
        summary=summary,
        chart_file=os.path.basename(chart_path)
    )

    html_path = os.path.join(out_dir, f"{base}.html")

    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_output)

    # =========================
    # PDF GENERATION (FIXED)
    # =========================
    pdf_path = os.path.join(out_dir, f"{base}.pdf")
    pdf_file = None

    try:
        pdf_template = env.get_template("report_pdf.html")

        pdf_html = pdf_template.render(
            app_name=safe_name,
            findings=mapped_findings,
            chart_base64=chart_base64,
            summary=summary,
            generated=datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        )

        with open(pdf_path, "wb") as f:
            pisa_status = pisa.CreatePDF(pdf_html, dest=f)

        if pisa_status.err:
            print("[PDF ERROR] PDF generation failed")
        else:
            pdf_file = os.path.basename(pdf_path)

    except Exception as e:
        print("[PDF EXCEPTION]", e)
        pdf_file = None

    # =========================
    # DOCX
    # =========================
    docx_path = None

    try:
        from docx import Document
        from docx.shared import Inches

        doc = Document()
        doc.add_heading(f"Security Report - {safe_name}", 0)

        doc.add_paragraph(f"Generated: {datetime.now()}")

        doc.add_picture(chart_path, width=Inches(5))

        table = doc.add_table(rows=1, cols=4)

        hdr = table.rows[0].cells
        hdr[0].text = "Issue"
        hdr[1].text = "Severity"
        hdr[2].text = "Description"
        hdr[3].text = "Recommendation"

        for fnd in mapped_findings:
            row = table.add_row().cells
            row[0].text = fnd.get("issue", "")
            row[1].text = fnd.get("severity", "")
            row[2].text = fnd.get("description", "")
            row[3].text = fnd.get("recommendation", "")

        docx_path = os.path.join(out_dir, f"{base}.docx")
        doc.save(docx_path)

    except Exception as e:
        print("[DOCX ERROR]", e)

    return {
        "html": html_path,
        "pdf": pdf_file,
        "docx": docx_path,
        "chart": chart_path
    }